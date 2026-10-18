"""Shared fixtures for the sidecar tests. Run with `npm run test:python`."""

import json
import os
import sqlite3
import subprocess
import sys
from pathlib import Path

import pyarrow as pa
import pyarrow.parquet as pq
import pytest

ROOT = Path(__file__).resolve().parents[2]
SIDECAR = ROOT / "python" / "sidecar"
RUNTIME = ROOT / "python" / "runtime"
# The bundled runtime, as the app runs it (not the test venv's interpreter).
RUNTIME_PYTHON = RUNTIME / "python.exe" if sys.platform == "win32" else RUNTIME / "bin" / "python3"
# Same flags as PYTHON_FLAGS in src/data.ts.
PYTHON_FLAGS = ["-E", "-s", "-B", "-u", "-X", "utf8"]

sys.path.insert(0, str(SIDECAR))
os.environ.setdefault("RUNTIME__DLTHUB_TELEMETRY", "false")

SALES = {
    "order_id": [1001, 1002, 1003, 1004],
    "order_date": ["2026-07-03", "2026-07-11", "2026-08-02", "2026-08-19"],
    "region": ["North", "South", "North", "West"],
    "revenue": [1250.5, 890.0, 2100.75, None],
}


def pytest_collection_modifyitems(config, items):
    if os.environ.get("PAWPAW_NETWORK_TESTS") == "1":
        return
    skip = pytest.mark.skip(reason="network test; set PAWPAW_NETWORK_TESTS=1 to run")
    for item in items:
        if "network" in item.keywords:
            item.add_marker(skip)


# Sample source files ----------------------------------------------------------------------------


def sales_rows(month: str = "") -> list[list]:
    """SALES as rows, optionally only orders from one month ("2026-07")."""
    rows = [list(row) for row in zip(*SALES.values())]
    return [row for row in rows if row[1].startswith(month)]


def write_csv(path: Path, header: list[str], rows: list[list]):
    import csv

    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(header)
        writer.writerows(["" if v is None else v for v in row] for row in rows)


@pytest.fixture
def sales_csv(tmp_path) -> Path:
    path = tmp_path / "Monthly Sales.csv"
    write_csv(path, list(SALES), sales_rows())
    return path


@pytest.fixture
def budget_xlsx(tmp_path) -> Path:
    """An Excel workbook with two sheets, Targets and Staff."""
    from openpyxl import Workbook

    path = tmp_path / "Budget.xlsx"
    workbook = Workbook()
    targets = workbook.active
    targets.title = "Targets"
    for row in (["region", "budget"], ["North", 5000], ["South", 4000]):
        targets.append(row)
    staff = workbook.create_sheet("Staff")
    for row in (["month", "headcount"], ["Jul", 12], ["Aug", 14], ["Sep", 15]):
        staff.append(row)
    workbook.save(path)
    return path


@pytest.fixture
def csv_folder(tmp_path) -> Path:
    """A folder of monthly CSV exports with the same columns."""
    folder = tmp_path / "monthly"
    folder.mkdir()
    for month in ("2026-07", "2026-08"):
        write_csv(folder / f"sales_{month}.csv", list(SALES), sales_rows(month))
    return folder


@pytest.fixture
def sqlite_db(tmp_path) -> Path:
    path = tmp_path / "shop.db"
    db = sqlite3.connect(path)
    db.executescript(
        """
        CREATE TABLE customers (id INTEGER PRIMARY KEY, name TEXT, city TEXT);
        INSERT INTO customers VALUES (1, 'Acme Pty Ltd', 'Brisbane'), (2, 'Bluegum Co', 'Hobart'),
                                     (3, 'Wattle & Sons', 'Perth');
        CREATE TABLE products (sku TEXT PRIMARY KEY, price REAL);
        INSERT INTO products VALUES ('A1', 9.5), ('B2', 12.0);
        """
    )
    db.commit()
    db.close()
    return path


PEOPLE = [{"id": 1, "name": "Ada"}, {"id": 2, "name": "Grace"}, {"id": 3, "name": "Linus"}]


@pytest.fixture
def json_array(tmp_path) -> Path:
    path = tmp_path / "people.json"
    path.write_text(json.dumps(PEOPLE), encoding="utf-8")
    return path


@pytest.fixture
def json_wrapper(tmp_path) -> Path:
    """An API-style response: the records are the one list inside a wrapper object."""
    path = tmp_path / "wrapped.json"
    path.write_text(json.dumps({"data": PEOPLE, "page": 1, "total": 3}), encoding="utf-8")
    return path


@pytest.fixture
def json_lines(tmp_path) -> Path:
    path = tmp_path / "events.jsonl"
    path.write_text("".join(json.dumps(p) + "\n" for p in PEOPLE) + "\n", encoding="utf-8")
    return path


@pytest.fixture
def json_nested(tmp_path) -> Path:
    """Records with a nested object (flattened into columns) and a list (a child table)."""
    orders = [
        {"id": 1, "customer": {"name": "Ada", "city": "Hobart"}, "items": [{"sku": "A1"}, {"sku": "B2"}]},
        {"id": 2, "customer": {"name": "Grace", "city": "Perth"}, "items": [{"sku": "A1"}]},
    ]
    path = tmp_path / "orders.json"
    path.write_text(json.dumps(orders), encoding="utf-8")
    return path


# A project data folder, as dlt would have written it ----------------------------------------------


def _write_raw(data_dir: Path, table: str, columns: dict, files: int = 1):
    """Write a table into <data_dir>/raw/<table>/ like dlt does, with dlt's bookkeeping columns."""
    folder = data_dir / "raw" / table
    folder.mkdir(parents=True, exist_ok=True)
    rows = len(next(iter(columns.values())))
    full = pa.table(
        {**columns, "_dlt_load_id": ["1792192294.07"] * rows, "_dlt_id": [f"id{i}" for i in range(rows)]}
    )
    step = -(-rows // files)
    for i in range(files):
        pq.write_table(full.slice(i * step, step), folder / f"1792192294.07.{i}.parquet")


@pytest.fixture
def project_data_dir(tmp_path) -> Path:
    """Where a project's data folder would be: nothing has been loaded into it yet."""
    return tmp_path / "project" / "data"


@pytest.fixture
def data_dir(project_data_dir) -> Path:
    """A project data folder with three loaded tables and dlt's own bookkeeping table.
    monthly_sales is split over two Parquet files, as dlt does for larger tables."""
    data = project_data_dir
    _write_raw(data, "monthly_sales", SALES, files=2)
    _write_raw(data, "budget_targets", {"region": ["North", "South"], "budget": [5000, 4000]})
    _write_raw(data, "budget_staff", {"month": ["Jul", "Aug", "Sep"], "headcount": [12, 14, 15]})
    loads = data / "raw" / "_dlt_loads"
    loads.mkdir()
    pq.write_table(pa.table({"load_id": ["1792192294.07"], "status": [0]}), loads / "loads.parquet")
    return data


def describe(data_dir: Path) -> dict:
    """Table name -> what describe_tables reports for it."""
    import datasets

    return {t["name"]: t for t in datasets.describe_tables({"data_dir": str(data_dir)})}


# Queries that must never succeed against the locked-down engine ------------------------------------

SYSTEM_FILE = "C:/Windows/win.ini" if sys.platform == "win32" else "/etc/hosts"

# (test id, SQL, where {outside} is the secret file, {outside_dir} its folder and {data} the
# data folder). Format them with escape_query().
ESCAPES = [
    ("read_system_file", f"SELECT * FROM read_csv('{SYSTEM_FILE}')"),
    ("read_text_outside_data_dir", "SELECT * FROM read_text('{outside}')"),
    ("read_csv_outside_data_dir", "SELECT * FROM read_csv('{outside}')"),
    ("read_blob_outside_data_dir", "SELECT * FROM read_blob('{outside}')"),
    ("from_first_syntax_outside_data_dir", "FROM read_csv('{outside}')"),
    ("dotdot_out_of_data_dir", "SELECT * FROM read_text('{data}/../secret.txt')"),
    ("sibling_folder_with_same_prefix", "SELECT * FROM read_text('{data}2/secret.txt')"),
    ("file_url_outside_data_dir", "SELECT * FROM read_text('file:///{outside}')"),
    ("glob_lists_files_outside", "SELECT * FROM glob('{outside_dir}/*')"),
    ("hidden_in_union", "SELECT 'x' AS a UNION ALL SELECT content FROM read_text('{outside}')"),
    ("hidden_in_cte", "WITH s AS (SELECT * FROM read_text('{outside}')) SELECT * FROM s"),
    ("web_request", "SELECT * FROM read_csv('https://jsonplaceholder.typicode.com/comments')"),
    ("set_external_access", "SET enable_external_access = true"),
    ("reset_setting", "RESET enable_external_access"),
    ("pragma_setting", "PRAGMA enable_external_access=true"),
    ("query_function_runs_set", "SELECT * FROM query('SET enable_external_access = true')"),
    ("copy_to_file", "COPY (SELECT 1) TO '{outside_dir}/x.csv'"),
    ("export_database", "EXPORT DATABASE '{outside_dir}/export'"),
    ("attach_database", "ATTACH '{outside_dir}/x.db'"),
    ("create_table", "CREATE TABLE x AS SELECT 1"),
    ("drop_view", "DROP VIEW monthly_sales"),
    ("two_statements", "SELECT 1; SELECT 2"),
    ("install_extension", "INSTALL httpfs"),
    ("load_extension", "LOAD httpfs"),
    ("call_procedure", "CALL pragma_version()"),
    ("syntax_error", "SELEC * FROM monthly_sales"),
    ("unknown_table", "SELECT * FROM nope"),
    ("empty_query", "   "),
]
ESCAPE_PARAMS = dict(argnames="sql", argvalues=[sql for _, sql in ESCAPES], ids=[name for name, _ in ESCAPES])


def escape_query(sql: str, data_dir: Path, outside: Path) -> str:
    """An ESCAPES query with its placeholders filled in."""
    return sql.format(outside=outside.as_posix(), outside_dir=outside.parent.as_posix(), data=data_dir.as_posix())


# A query that never finishes within a test's time limit. A cross join rather than one huge
# range(), which DuckDB could one day answer from the range's size without scanning it.
SLOW = "SELECT count(*) FROM range(1000000) a, range(1000000) b"


# Running runner.py like the app does ----------------------------------------------------------------


@pytest.fixture
def run_job(tmp_path):
    """Runs one job through python/sidecar/runner.py in the bundled runtime, the same way
    src/data.ts does, and returns its parsed events. dlt's state lives in a folder per test."""
    env = {
        **os.environ,
        "RUNTIME__DLTHUB_TELEMETRY": "false",
        "DLT_DATA_DIR": str(tmp_path / "dlt-state"),
    }

    def run(job: dict | str, timeout: float = 300) -> list[dict]:
        line = job if isinstance(job, str) else json.dumps(job)
        result = subprocess.run(
            [str(RUNTIME_PYTHON), *PYTHON_FLAGS, str(SIDECAR / "runner.py")],
            input=(line + "\n").encode("utf-8"),
            capture_output=True,
            env=env,
            timeout=timeout,
        )
        events = [json.loads(line) for line in result.stdout.decode("utf-8").splitlines() if line.strip()]
        assert events, f"runner.py wrote no events. stderr:\n{result.stderr.decode('utf-8', 'replace')}"
        # Exit status matches the outcome.
        assert (result.returncode == 0) == (events[-1]["event"] == "result"), events[-1]
        return events

    return run


def result_of(events: list[dict]):
    """The result of a successful job, with a helpful failure message otherwise."""
    last = events[-1]
    assert last["event"] == "result", f"{last.get('message')}\n{last.get('detail', '')}"
    return last["result"]


def error_of(events: list[dict]) -> str:
    last = events[-1]
    assert last["event"] == "error", f"expected an error, got {last}"
    return last["message"]


@pytest.fixture
def outside(tmp_path) -> Path:
    """A file next to the project, outside its data folder."""
    path = tmp_path / "secret.txt"
    path.write_text("secret\nhunter2\n", encoding="utf-8")
    return path


@pytest.fixture
def sibling_secret(data_dir) -> Path:
    """A secret file in a sibling folder (<data_dir>2) whose name starts with the data folder's."""
    sibling = data_dir.parent / f"{data_dir.name}2"
    sibling.mkdir()
    path = sibling / "secret.txt"
    path.write_text("secret\nhunter2\n", encoding="utf-8")
    return path
