"""Helpers of the file connectors (python/sidecar/connectors/files.py)."""

import io
import json

import pytest

from connectors.files import _column_names, _detect_format, _json_records, _read_csv, _read_excel, _read_json


@pytest.mark.parametrize(
    ("url", "content_type", "content", "expected"),
    [
        # The URL's extension wins.
        ("https://x.test/data.csv", "application/json", b"[1]", "csv"),
        ("https://x.test/data.TSV", "", b"", "csv"),
        ("https://x.test/data.txt", "", b"", "csv"),
        ("https://x.test/data.json?token=abc", "text/csv", b"", "json"),
        ("https://x.test/data.jsonl", "", b"", "jsonl"),
        ("https://x.test/data.ndjson", "", b"", "jsonl"),
        ("https://x.test/data.parquet", "", b"", "parquet"),
        ("https://x.test/data.xlsx", "", b"", "xlsx"),
        ("https://x.test/data.xlsm", "", b"", "xlsx"),
        # Then the Content-Type.
        ("https://x.test/api/items", "application/json; charset=utf-8", b"", "json"),
        ("https://x.test/api/items", "application/x-ndjson", b"", "jsonl"),
        ("https://x.test/api/items", "application/jsonl", b"", "jsonl"),
        ("https://x.test/export", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", b"", "xlsx"),
        ("https://x.test/export", "application/vnd.ms-excel", b"", "xlsx"),
        ("https://x.test/export", "application/vnd.apache.parquet", b"", "parquet"),
        # Then the content itself.
        ("https://x.test/download", "application/octet-stream", b"PAR1\x00\x01", "parquet"),
        ("https://x.test/download", "application/octet-stream", b"PK\x03\x04", "xlsx"),
        ("https://x.test/download", "text/plain", b'[{"a": 1}]', "json"),
        ("https://x.test/download", "text/plain", b'{"data": []}', "json"),
        ("https://x.test/download", "text/plain", b'\xef\xbb\xbf \r\n\t[{"a": 1}]', "json"),
        # Otherwise CSV.
        ("https://x.test/download", "text/plain", b"a,b\n1,2\n", "csv"),
        ("https://x.test/download", "", b"", "csv"),
        ("https://x.test/file.unknown", "", b"a;b\n", "csv"),
    ],
)
def test_detect_format(url, content_type, content, expected):
    assert _detect_format(url, content_type, content) == expected


def test_json_records_from_list():
    assert _json_records([{"a": 1}, {"a": 2}]) == [{"a": 1}, {"a": 2}]


def test_json_records_from_wrapper_object():
    assert _json_records({"data": [{"a": 1}], "page": 1, "next": None}) == [{"a": 1}]


def test_json_records_from_single_object():
    assert _json_records({"a": 1, "b": {"c": 2}}) == [{"a": 1, "b": {"c": 2}}]


def test_json_records_object_with_several_lists_is_one_record():
    data = {"users": [{"a": 1}], "groups": [{"b": 2}]}
    assert _json_records(data) == [data]


def test_json_records_from_scalar():
    assert _json_records(42) == [{"value": 42}]


def _all(chunks) -> list:
    return [record for chunk in chunks for record in chunk]


def test_read_json_array_from_path(json_array):
    assert _all(_read_json(json_array)) == json.loads(json_array.read_text(encoding="utf-8"))


def test_read_json_wrapper_from_path(json_wrapper):
    assert [r["name"] for r in _all(_read_json(json_wrapper))] == ["Ada", "Grace", "Linus"]


def test_read_json_lines(json_lines):
    assert [r["id"] for r in _all(_read_json(json_lines))] == [1, 2, 3]


def test_read_json_from_bytes_with_bom():
    content = io.BytesIO('﻿[{"name": "Zoë"}]'.encode("utf-8"))
    assert _all(_read_json(content)) == [{"name": "Zoë"}]


def test_read_json_lines_from_bytes():
    content = io.BytesIO(b'{"a": 1}\r\n\r\n{"a": 2}\r\n')
    assert _all(_read_json(content)) == [{"a": 1}, {"a": 2}]


def test_read_json_chunks_large_files(monkeypatch):
    import connectors.files as files

    monkeypatch.setattr(files, "CSV_CHUNK_ROWS", 2)
    content = io.BytesIO(json.dumps([{"i": i} for i in range(5)]).encode())
    assert [len(chunk) for chunk in _read_json(content)] == [2, 2, 1]


def test_read_json_rejects_invalid_json():
    with pytest.raises(json.JSONDecodeError):
        _all(_read_json(io.BytesIO(b"{not json")))


# CSV (DuckDB) -------------------------------------------------------------------------------------


def _csv_table(source, **options):
    import pyarrow as pa

    return pa.Table.from_batches(list(_read_csv(source, **options)))


def test_read_csv_detects_types(sales_csv):
    table = _csv_table(sales_csv)
    assert table.column_names == ["order_id", "order_date", "region", "revenue"]
    assert [str(t) for t in table.schema.types] == ["int64", "date32[day]", "string", "double"]
    assert table.num_rows == 4


def test_read_csv_detects_semicolons_from_bytes():
    table = _csv_table(io.BytesIO("name;amount\nAcme;1,5\nBluegum;2\n".encode()))
    assert table.column_names == ["name", "amount"]
    assert table.num_rows == 2


def test_read_csv_explicit_tab_delimiter(tmp_path):
    path = tmp_path / "data.tsv"
    path.write_text("a\tb\n1\tx\n", encoding="utf-8")
    assert _csv_table(path, delimiter="\t").to_pylist() == [{"a": 1, "b": "x"}]


# Excel (openpyxl) ---------------------------------------------------------------------------------


def test_column_names_fill_blanks_and_dedupe():
    assert _column_names(["Region", None, "  ", "Region", "Region", 2026]) == [
        "Region",
        "column_2",
        "column_3",
        "Region_2",
        "Region_3",
        "2026",
    ]


def test_read_excel_sheet(budget_xlsx):
    assert _all(_read_excel(budget_xlsx, "Staff")) == [
        {"month": "Jul", "headcount": 12},
        {"month": "Aug", "headcount": 14},
        {"month": "Sep", "headcount": 15},
    ]


def test_read_excel_defaults_to_first_sheet_and_skips_empty_rows(tmp_path):
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    for row in (["name", None, "name"], ["a", 1, "b"], [None, None, None], ["c", 2, "d"]):
        sheet.append(row)
    path = tmp_path / "book.xlsx"
    workbook.save(path)
    assert _all(_read_excel(io.BytesIO(path.read_bytes()))) == [
        {"name": "a", "column_2": 1, "name_2": "b"},
        {"name": "c", "column_2": 2, "name_2": "d"},
    ]


def test_read_excel_formulas_give_values(tmp_path):
    from openpyxl import Workbook, load_workbook

    path = tmp_path / "formulas.xlsx"
    workbook = Workbook()
    workbook.active.append(["a", "b"])
    workbook.active.append([2, "=A2*2"])
    workbook.save(path)
    # openpyxl doesn't calculate formulas, and a file it wrote has no cached values, so the cell
    # reads as empty; files saved by Excel carry the calculated value.
    assert _all(_read_excel(path)) == [{"a": 2, "b": None}]
    assert load_workbook(path).active["B2"].value == "=A2*2"


def test_file_reader_rejects_unknown_format():
    from connectors.files import file_reader

    with pytest.raises(ValueError, match="Unsupported format: xml"):
        file_reader("xml")
