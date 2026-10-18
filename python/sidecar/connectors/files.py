"""Connectors for files on this computer or on the web. Bundled with the app: no downloads needed."""

import io
import json
from pathlib import Path

import dlt

from .base import FILES, Connector, Field, Option, lines, required, text

DATA_FILE_FILTERS = [
    ("Data files", ["csv", "tsv", "txt", "json", "jsonl", "ndjson", "parquet", "xlsx", "xlsm"]),
    ("All files", ["*"]),
]
FORMAT_OPTIONS = [
    Option("csv", "CSV"),
    Option("jsonl", "JSON Lines"),
    Option("parquet", "Parquet"),
]
CSV_CHUNK_ROWS = 50_000


def _read_csv(source, delimiter: str | None = None):
    """CSV rows as Arrow batches, from a path or a file object. DuckDB detects the delimiter,
    header and column types (dates, numbers, …) from the data."""
    import duckdb

    con = duckdb.connect()
    try:
        options = {"delimiter": delimiter} if delimiter else {}
        relation = con.read_csv(str(source) if isinstance(source, Path) else source, **options)
        yield from relation.to_arrow_reader(CSV_CHUNK_ROWS)
    finally:
        con.close()


def _column_names(header) -> list[str]:
    """Names from a spreadsheet's first row: blanks get column_<n>, duplicates get a suffix."""
    names: list[str] = []
    for i, value in enumerate(header, 1):
        name = str(value).strip() if value is not None else ""
        name = name or f"column_{i}"
        base, n = name, 2
        while name in names:
            name, n = f"{base}_{n}", n + 1
        names.append(name)
    return names


def _excel_sheet_names(source) -> list[str]:
    from openpyxl import load_workbook

    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        return workbook.sheetnames
    finally:
        workbook.close()


def _read_excel(source, sheet: str | None = None):
    """A sheet's rows as dicts, using its first row as column names. Formulas give their last
    calculated values. Defaults to the first sheet."""
    from openpyxl import load_workbook

    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        rows = workbook[sheet or workbook.sheetnames[0]].iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            return
        names = _column_names(header)
        batch = []
        for row in rows:
            if all(value is None for value in row):
                continue
            batch.append(dict(zip(names, row)))
            if len(batch) >= CSV_CHUNK_ROWS:
                yield batch
                batch = []
        if batch:
            yield batch
    finally:
        workbook.close()


def _json_records(data) -> list:
    """The records in a parsed JSON document: a top-level list, or the one list inside a
    wrapper object such as {"data": [...]}; otherwise the object itself is a single record."""
    if isinstance(data, list):
        return data
    if isinstance(data, dict):
        lists = [value for value in data.values() if isinstance(value, list)]
        return lists[0] if len(lists) == 1 else [data]
    return [{"value": data}]


def _read_json(source):
    """JSON or JSON Lines, from a path or bytes. Records go to dlt as dicts, so nested objects
    and lists become their own columns and child tables."""
    raw = source.read_bytes() if isinstance(source, Path) else source.getvalue()
    content = raw.decode("utf-8-sig")
    try:
        records = _json_records(json.loads(content))
    except json.JSONDecodeError:
        # Not a single document: JSON Lines, one record per line.
        records = [json.loads(line) for line in content.splitlines() if line.strip()]
    for start in range(0, len(records), CSV_CHUNK_ROWS):
        yield records[start : start + CSV_CHUNK_ROWS]


def _read_parquet(path: Path):
    import pyarrow.parquet as pq

    yield from pq.ParquetFile(path).iter_batches(batch_size=CSV_CHUNK_ROWS)


def _frames(name: str, read):
    """A resource that yields the batches from `read()` (Arrow batches or lists of dicts) when
    loading, not before."""

    @dlt.resource(name=name)
    def frames():
        yield from read()

    return frames


def _file_resources(path: Path, used_names: set[str]) -> list:
    if not path.is_file():
        raise ValueError(f"File not found: {path}")
    ext = path.suffix.lower()

    def unique(name: str) -> str:
        candidate, n = name, 2
        while candidate in used_names:
            candidate, n = f"{name}_{n}", n + 1
        used_names.add(candidate)
        return candidate

    if ext in (".xlsx", ".xlsm"):
        sheets = _excel_sheet_names(path)
        resources = []
        for sheet in sheets:
            name = unique(path.stem if len(sheets) == 1 else f"{path.stem} {sheet}")
            resources.append(_frames(name, lambda sheet=sheet: _read_excel(path, sheet)))
        return resources
    if ext in (".csv", ".tsv", ".txt"):
        return [_frames(unique(path.stem), lambda: _read_csv(path, "\t" if ext == ".tsv" else None))]
    if ext in (".json", ".jsonl", ".ndjson"):
        return [_frames(unique(path.stem), lambda: _read_json(path))]
    if ext == ".parquet":
        return [_frames(unique(path.stem), lambda: _read_parquet(path))]
    raise ValueError(f"Unsupported file type: {path.name}")


def _build_local_files(settings: dict):
    paths = lines(settings, "files")
    if not paths:
        raise ValueError("Choose at least one file.")
    used: set[str] = set()
    return [r for p in paths for r in _file_resources(Path(p), used)]


def _build_folder(settings: dict):
    from dlt.sources.filesystem import filesystem

    folder = Path(required(settings, "folder", "Folder"))
    if not folder.is_dir():
        raise ValueError(f"Folder not found: {folder}")
    fmt = text(settings, "format", "csv")
    files = filesystem(bucket_url=folder.resolve().as_uri(), file_glob=text(settings, "pattern", f"**/*.{fmt}"))
    return (files | file_reader(fmt)).with_name(text(settings, "table", folder.name))


def file_reader(fmt: str):
    """dlt's reader for files listed by its filesystem source. CSV goes through DuckDB, which
    detects types; the other readers don't need pandas either."""
    from dlt.sources.filesystem import read_csv_duckdb, read_jsonl, read_parquet

    if fmt == "csv":
        return read_csv_duckdb(chunk_size=CSV_CHUNK_ROWS, use_pyarrow=True)
    readers = {"jsonl": read_jsonl, "parquet": read_parquet}
    if fmt not in readers:
        raise ValueError(f"Unsupported format: {fmt}")
    return readers[fmt]()


def _build_url(settings: dict):
    from dlt.sources.helpers import requests

    url = required(settings, "url", "Web address")
    chosen_format = text(settings, "format")
    table = text(settings, "table") or Path(url.split("?")[0]).stem or "download"

    def read():
        response = requests.get(url, timeout=120)
        response.raise_for_status()
        fmt = chosen_format or _detect_format(url, response.headers.get("Content-Type", ""), response.content)
        content = io.BytesIO(response.content)
        if fmt == "csv":
            yield from _read_csv(content)
        elif fmt in ("jsonl", "json"):
            yield from _read_json(content)
        elif fmt == "parquet":
            import pyarrow.parquet as pq

            yield pq.read_table(content)
        elif fmt == "xlsx":
            yield from _read_excel(content)
        else:
            raise ValueError(f"Unsupported format: {fmt}")

    return _frames(table, read)


EXTENSION_FORMATS = {
    "csv": "csv", "tsv": "csv", "txt": "csv",
    "json": "json", "jsonl": "jsonl", "ndjson": "jsonl",
    "parquet": "parquet",
    "xlsx": "xlsx", "xlsm": "xlsx",
}


def _detect_format(url: str, content_type: str, content: bytes) -> str:
    """The file format from the URL's extension, then the server's Content-Type, then the
    content itself. Many APIs and download links have no extension at all."""
    ext = Path(url.split("?")[0]).suffix.lower().lstrip(".")
    if ext in EXTENSION_FORMATS:
        return EXTENSION_FORMATS[ext]

    content_type = content_type.lower()
    if "ndjson" in content_type or "jsonl" in content_type:
        return "jsonl"
    if "json" in content_type:
        return "json"
    if "spreadsheetml" in content_type or "ms-excel" in content_type:
        return "xlsx"
    if "parquet" in content_type:
        return "parquet"

    if content.startswith(b"PAR1"):
        return "parquet"
    if content.startswith(b"PK"):  # A ZIP container, as .xlsx files are.
        return "xlsx"
    if content.lstrip(b"\xef\xbb\xbf \t\r\n")[:1] in (b"[", b"{"):
        return "json"
    return "csv"


CONNECTORS: list[Connector] = [
    Connector(
        id="local_files",
        name="Files on this computer",
        category=FILES,
        description="CSV, Excel, JSON and Parquet files. Each file (or Excel sheet) becomes a table.",
        keywords=["csv", "excel", "xlsx", "spreadsheet", "json", "parquet", "upload"],
        fields=[
            Field("files", "Files", type="files", required=True, file_filters=DATA_FILE_FILTERS),
        ],
        build=_build_local_files,
    ),
    Connector(
        id="folder",
        name="Folder of files",
        category=FILES,
        description="Combine every matching file in a folder, such as monthly exports, into one table.",
        docs_url="https://dlthub.com/docs/dlt-ecosystem/verified-sources/filesystem",
        keywords=["directory", "csv", "exports", "combine"],
        fields=[
            Field("folder", "Folder", type="folder", required=True),
            Field("format", "File format", type="select", default="csv", options=FORMAT_OPTIONS),
            Field("table", "Table name", help="Defaults to the folder name."),
            Field(
                "pattern",
                "Files to include",
                placeholder="**/*.csv",
                advanced=True,
                help="A pattern such as sales_*.csv. ** includes subfolders. Defaults to every file of the chosen format.",
            ),
        ],
        build=_build_folder,
    ),
    Connector(
        id="url",
        name="File from a web address",
        category=FILES,
        description="Download a CSV, Excel, JSON or Parquet file from a URL.",
        keywords=["http", "https", "download", "link", "csv"],
        fields=[
            Field("url", "Web address", required=True, placeholder="https://example.com/data.csv"),
            Field("table", "Table name", help="Defaults to the file name in the address."),
            Field(
                "format",
                "File format",
                type="select",
                advanced=True,
                help="Detected from the address when left blank.",
                options=[Option("", "Detect automatically"), *FORMAT_OPTIONS, Option("json", "JSON"), Option("xlsx", "Excel")],
            ),
        ],
        build=_build_url,
    ),
]
